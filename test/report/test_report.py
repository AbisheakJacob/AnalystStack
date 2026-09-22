"""Tests for AnalystStack.report.report."""

from typing import Any

import pandas as pd
import pytest
from openpyxl import load_workbook

from AnalystStack.report.report import DataQualityReport, ReportSection


@pytest.fixture
def report() -> DataQualityReport:
    return DataQualityReport(
        title="Orders Data Quality Report",
        generated_at=pd.Timestamp("2026-01-01 12:00", tz="UTC"),
        sections=[
            ReportSection(name="Tables", dataframe=pd.DataFrame({"table_name": ["orders"]})),
            ReportSection(name="Notes", dataframe=None, narrative="Nothing else to report."),
        ],
    )


def test_write_markdown_includes_title_and_sections(report, tmp_path):
    file_path = tmp_path / "report.md"
    report.write(str(file_path))

    content = file_path.read_text(encoding="utf-8")
    assert content.startswith("# Orders Data Quality Report")
    assert "## Tables" in content
    assert "## Notes" in content
    assert "Nothing else to report." in content
    assert "orders" in content


def test_write_txt_includes_title_and_sections(report, tmp_path):
    file_path = tmp_path / "report.txt"
    report.write(str(file_path), format="txt")

    content = file_path.read_text(encoding="utf-8")
    assert content.startswith("Orders Data Quality Report")
    assert "Tables" in content
    assert "Nothing else to report." in content


def test_write_excel_creates_one_sheet_per_dataframe_section(report, tmp_path):
    file_path = tmp_path / "report.xlsx"
    report.write(str(file_path), format="excel")

    workbook = load_workbook(str(file_path))
    assert "Tables" in workbook.sheetnames
    assert "Notes" not in workbook.sheetnames  # no dataframe -> no sheet
    rows = list(workbook["Tables"].iter_rows(values_only=True))
    assert rows == [("table_name",), ("orders",)]


def test_write_excel_truncates_long_sheet_names(tmp_path):
    long_name = "A" * 50
    report = DataQualityReport(
        title="t",
        generated_at=pd.Timestamp.now(tz="UTC"),
        sections=[ReportSection(name=long_name, dataframe=pd.DataFrame({"a": [1]}))],
    )
    file_path = tmp_path / "report.xlsx"
    report.write(str(file_path), format="excel")

    workbook = load_workbook(str(file_path))
    assert long_name[:31] in workbook.sheetnames


def test_write_rejects_unsupported_format(report, tmp_path):
    bogus_format: Any = "pdf"
    with pytest.raises(ValueError, match="Unsupported format"):
        report.write(str(tmp_path / "report.pdf"), format=bogus_format)
